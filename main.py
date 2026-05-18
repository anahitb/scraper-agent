#!/usr/bin/env python3
"""
Bulk Internet Agreement Scraper Agent

Searches for competitor bulk internet/telecommunications service agreements,
downloads PDFs, verifies they are genuine bulk agreements, and summarizes key terms.

Usage:
    python main.py                          # uses PROVIDERS list from config.py
    python main.py "Comcast" "AT&T"        # override with specific providers
"""

import argparse
import json
import logging
import os
import sys
import time
from tqdm import tqdm
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from config import PROVIDERS, OUTPUT_DIR, REPORT_FILE
from searcher import search_provider
from url_filter import filter_results
from pdf_handler import process_url
from analyzer import analyze

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    handlers=[
        logging.StreamHandler(sys.stdout),
        logging.FileHandler("output/scraper.log"),
    ],
)
logger = logging.getLogger(__name__)

EXCEL_FILE = "output/bulk_agreements.xlsx"

COLUMNS = [
    ("Provider (ISP)",           lambda a, kt, r: a.get("provider") or r.get("provider")),
    ("Property Name",            lambda a, kt, r: a.get("property_name")),
    ("Property Address",         lambda a, kt, r: a.get("property_address")),
    ("Number of Units",          lambda a, kt, r: kt.get("number_of_units")),
    ("Service Commitment Period",lambda a, kt, r: kt.get("service_commitment_period")),
    ("Auto Renewal",             lambda a, kt, r: kt.get("auto_renewal")),
    ("Monthly Fee / Unit",       lambda a, kt, r: kt.get("monthly_fee_per_unit")),
    ("Total Monthly Billing",    lambda a, kt, r: kt.get("total_monthly_billing")),
    ("Door Fee / Unit",          lambda a, kt, r: kt.get("door_fee_per_unit")),
    ("Internet Speed",           lambda a, kt, r: kt.get("internet_speed")),
    ("Services Covered",         lambda a, kt, r: ", ".join(kt.get("services_covered") or [])),
    ("Marketing Exclusivity",    lambda a, kt, r: kt.get("marketing_exclusivity")),
    ("Annual Rate Increase Cap", lambda a, kt, r: kt.get("annual_rate_increase_cap")),
    ("Capital Investment",       lambda a, kt, r: kt.get("capital_investment_by_operator")),
    ("Notable Clauses",          lambda a, kt, r: "; ".join(kt.get("notable_clauses") or [])),
    ("Confidence",               lambda a, kt, r: a.get("confidence")),
    ("Source URL",               lambda a, kt, r: r.get("url")),
    ("Local PDF",                lambda a, kt, r: r.get("local_path")),
]


def save_excel(results: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bulk Agreements"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    alt_fill = PatternFill("solid", fgColor="D6E4F0")

    # Header row
    for col_idx, (header, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    ws.row_dimensions[1].height = 30

    # Data rows
    for row_idx, r in enumerate(results, start=2):
        a = r.get("analysis", {})
        kt = a.get("key_takeaways", {})
        fill = alt_fill if row_idx % 2 == 0 else PatternFill()
        for col_idx, (_, extractor) in enumerate(COLUMNS, start=1):
            try:
                value = extractor(a, kt, r)
            except Exception:
                value = ""
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Auto-width columns
    for col_idx in range(1, len(COLUMNS) + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        max_len = max(
            (len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1)),
            default=10,
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 50)

    ws.freeze_panes = "A2"
    wb.save(EXCEL_FILE)
    logger.info(f"Excel report saved to {EXCEL_FILE}")


def run(providers: list[str]) -> list[dict]:
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    all_results = []

    for provider in providers:
        logger.info(f"=== Processing provider: {provider} ===")

        raw_results = search_provider(provider)
        filtered = filter_results(raw_results)

        if not filtered:
            logger.info(f"No relevant URLs found for {provider}")
            continue

        logger.info(f"Processing {len(filtered)} candidate URLs for {provider}...")
        processed = []
        for result in tqdm(filtered, desc=f"{provider} — fetching"):
            processed.append(process_url(result))
            time.sleep(0.5)

        downloaded = [r for r in processed if r.get("status") == "downloaded"]
        logger.info(f"{provider}: {len(downloaded)}/{len(filtered)} URLs passed content verification")

        analyzed = []
        for result in tqdm(downloaded, desc=f"{provider} — analyzing"):
            analyzed.append(analyze(result))

        confirmed = [
            r for r in analyzed
            if r.get("analysis", {}).get("is_bulk_agreement") is True
        ]
        logger.info(f"{provider}: {len(confirmed)} confirmed bulk agreements")
        all_results.extend(confirmed)

    return all_results


def save_report(results: list[dict]) -> None:
    report = []
    for r in results:
        entry = {k: v for k, v in r.items() if k not in ("full_text", "pdf_bytes")}
        report.append(entry)

    with open(REPORT_FILE, "w") as f:
        json.dump(report, f, indent=2)
    logger.info(f"JSON report saved to {REPORT_FILE}")

    save_excel(report)

    print("\n" + "=" * 60)
    print(f"RESULTS: {len(report)} confirmed bulk agreements found")
    print(f"Excel:   {EXCEL_FILE}")
    print(f"JSON:    {REPORT_FILE}")
    print("=" * 60)
    for r in report:
        a = r.get("analysis", {})
        kt = a.get("key_takeaways", {})
        print(f"\nProvider : {a.get('provider', r.get('provider'))}")
        print(f"Property : {a.get('property_name', 'unknown')}")
        print(f"Units    : {kt.get('number_of_units', 'unknown')}")
        print(f"URL      : {r['url']}")
        print(f"Fee/Unit : {kt.get('monthly_fee_per_unit', 'unknown')}")
        print(f"Term     : {kt.get('service_commitment_period', 'unknown')}")
        print(f"Exclusivity: {kt.get('marketing_exclusivity', 'unknown')}")
        print("-" * 40)


def main():
    parser = argparse.ArgumentParser(description="Bulk Internet Agreement Scraper")
    parser.add_argument(
        "providers",
        nargs="*",
        help="Provider names to search (defaults to config.py PROVIDERS list)",
    )
    args = parser.parse_args()

    providers = args.providers if args.providers else PROVIDERS
    logger.info(f"Starting scraper for {len(providers)} providers: {providers}")

    results = run(providers)
    save_report(results)


if __name__ == "__main__":
    main()
